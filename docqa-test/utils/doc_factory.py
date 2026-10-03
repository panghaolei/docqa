"""测试语料工厂：按需生成 txt / md / docx / pdf 文档。

为什么要在测试里"生成"文档而不是放几个固定文件？
- 知识库内容必须可控：评测集的问题、期望答案、期望来源都依赖语料里的确定事实；
- 覆盖多种解析格式（Tika 对 PDF/Word 的解析质量是高风险点）；
- 便于构造边界语料（超长文本、空文件、二进制文件）。
"""

from __future__ import annotations

import os
import tempfile
from pathlib import Path

# =====================================================================
# 知识库语料：一份虚构的《员工手册》，内置大量确定事实，供 RAG 评测使用
# =====================================================================
KB_DOC_NAME = "星云科技员工手册_测试语料.txt"

KB_TEXT = """星云科技《员工手册（2026 版）》—— 内部培训测试资料

一、工作时间与考勤
公司实行标准工时制，工作时间为每周一至周五 9:00 至 18:00，午休时间为 12:00 至 13:30。
全体员工需通过企业考勤系统打卡，迟到累计超过 3 次将影响当月绩效奖金。
研发岗位可申请弹性工作制，最晚 10:30 前到岗即可，但每日在岗时长不得少于 8 小时。
连续远程办公超过 5 个工作日，需由部门负责人与人力资源部共同审批。

二、假期制度
1. 年假：入职满 1 年可享受 5 天年假；满 3 年可享受 10 天；满 5 年及以上可享受 15 天。
2. 病假：需提供二级及以上医院证明，病假期间按基本工资的 80% 发放。
3. 事假：按日扣减当月工资，全年累计事假不得超过 15 天。
4. 婚假 10 天，产假 158 天，陪产假 15 天，产检假每次半天且不计入病假。

三、加班与调休
工作日加班按 1.5 倍计算加班费，休息日加班按 2 倍计算，法定节假日加班按 3 倍计算。
加班需提前在 OA 系统提交申请，经直属主管审批后方可生效。加班满 4 小时可申请调休半天。
工作日加班至 21:00 之后可申请 30 元餐补，22:30 之后可实报实销打车费用。
调休需在加班发生后 60 天内使用完毕，逾期自动作废。

四、费用报销
差旅费需在出差结束后 7 个工作日内提交报销单，由部门经理审批。
单笔报销金额超过 5000 元时，还需财务总监审批。
市内交通费每月上限 300 元，报销时需附发票原件。
一线城市住宿标准为每晚 500 元，其他城市为每晚 350 元；高铁一律乘坐二等座。

五、试用期与转正
新员工试用期为 3 个月，试用期工资为转正后工资的 80%。
转正考核在试用期结束前 10 个工作日启动，考核结果分为优秀、合格、不合格三档。
试用期内提前完成关键项目并获部门负责人推荐的员工，可申请提前 1 个月转正。

六、薪酬与调薪
公司每年进行两次调薪评估，时间分别在 4 月和 10 月。
年度绩效为 A 的员工调薪幅度不低于 15%，绩效为 C 的员工不参与本次调薪。
公司实行十三薪制度，于每年春节前一周发放；绩效系数区间为 0.8 至 1.5。

七、培训与发展
每位员工每年享有 3000 元的培训预算，可用于外部课程、认证考试与书籍采购。
员工每年至少完成 16 学时的内部培训课程。
取得与岗位相关的行业认证后，可按考试费用的 70% 报销，每人每年上限 5000 元。

八、保密与竞业限制
员工须对公司的技术资料、客户名单与财务数据严格保密。
核心技术岗位员工离职后 2 年内不得从事与公司业务直接竞争的工作，公司按月支付竞业补偿金。
客户名单、算法模型与财务数据属于公司核心机密，任何外发行为均需信息安全部审批。
违反保密义务的，公司有权解除劳动合同并追究法律责任。
"""

# 评测集使用的"确定事实"，与 evaluation/dataset.yaml 保持一致
KB_FACTS = {
    "annual_leave_3y": "10 天",
    "annual_leave_5y": "15 天",
    "ot_weekday": "1.5 倍",
    "ot_holiday": "3 倍",
    "reimburse_submit": "7 个工作日",
    "reimburse_approve": "5000 元",
    "probation": "3 个月",
    "salary_review": "4 月和 10 月",
    "training_budget": "3000 元",
    "non_compete": "2 年",
}


# =====================================================================
# 各格式文档生成
# =====================================================================
def make_txt(dest: Path, content: str = KB_TEXT) -> Path:
    dest.parent.mkdir(parents=True, exist_ok=True)
    dest.write_text(content, encoding="utf-8")
    return dest


def make_md(dest: Path, content: str = KB_TEXT) -> Path:
    """把纯文本转成 Markdown（加标题层级），用于覆盖 md 解析分支。"""
    lines = []
    for line in content.splitlines():
        if line.startswith("一、") or line.startswith("二、") or line.startswith("三、") \
                or line.startswith("四、") or line.startswith("五、") or line.startswith("六、") \
                or line.startswith("七、") or line.startswith("八、"):  # noqa: PLR1714
            lines.append(f"## {line}")
        else:
            lines.append(line)
    dest.parent.mkdir(parents=True, exist_ok=True)
    dest.write_text("\n".join(lines), encoding="utf-8")
    return dest


def make_docx(dest: Path, content: str = KB_TEXT) -> Path:
    """生成 Word 文档，覆盖 Tika 的 OOXML 解析分支。"""
    from docx import Document  # 延迟导入，未安装 python-docx 时其它用例不受影响

    doc = Document()
    doc.add_heading("星云科技员工手册（2026 版）", level=1)
    for line in content.splitlines():
        line = line.strip()
        if not line:
            continue
        doc.add_paragraph(line)
    dest.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(dest))
    return dest


def make_pdf(dest: Path, content: str = KB_TEXT) -> Path:
    """生成 PDF 文档（注册中文字体 STSong-Light，避免中文乱码/方框）。"""
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.pdfgen import canvas

    font_name = "STSong-Light"
    try:
        pdfmetrics.registerFont(UnicodeCIDFont(font_name))
    except Exception:  # pragma: no cover - 字体注册失败时退化为英文，不影响测试流程
        font_name = "Helvetica"

    dest.parent.mkdir(parents=True, exist_ok=True)
    c = canvas.Canvas(str(dest), pagesize=A4)
    width, height = A4
    c.setFont(font_name, 11)
    y = height - 60
    for line in content.splitlines():
        if y < 60:
            c.showPage()
            c.setFont(font_name, 11)
            y = height - 60
        c.drawString(50, y, line)
        y -= 18
    c.save()
    return dest


def make_html(dest: Path, content: str = KB_TEXT) -> Path:
    html = ("<html><head><meta charset='utf-8'><title>员工手册</title></head><body>"
            + "".join(f"<p>{line}</p>" for line in content.splitlines() if line.strip())
            + "</body></html>")
    dest.parent.mkdir(parents=True, exist_ok=True)
    dest.write_text(html, encoding="utf-8")
    return dest


# =====================================================================
# 边界 / 异常语料
# =====================================================================
def make_empty_file(dest: Path) -> Path:
    dest.parent.mkdir(parents=True, exist_ok=True)
    dest.write_bytes(b"")
    return dest


def make_fake_binary(dest: Path, size: int = 2048) -> Path:
    """伪造一个非文档二进制文件（.exe），验证后端对非法格式的处理。"""
    dest.parent.mkdir(parents=True, exist_ok=True)
    dest.write_bytes(os.urandom(size))
    return dest


def make_long_text_doc(dest: Path, repeat: int = 12) -> Path:
    """生成超长文档（约 1.2 万字），用于验证分块数量与截断逻辑。"""
    parts = [f"第 {i + 1} 章节\n{KB_TEXT}" for i in range(repeat)]
    dest.parent.mkdir(parents=True, exist_ok=True)
    dest.write_text("\n\n".join(parts), encoding="utf-8")
    return dest


def make_large_padded_doc(dest: Path, target_kb: int = 600) -> Path:
    """生成约 target_kb KB 的文档，用于上传性能与大小边界测试。"""
    block = KB_TEXT + "\n"
    body = block * max(1, (target_kb * 1024) // len(block.encode("utf-8")))
    dest.parent.mkdir(parents=True, exist_ok=True)
    dest.write_text(body, encoding="utf-8")
    return dest


# =====================================================================
# 便捷入口
# =====================================================================
FACTORY = {
    "txt": make_txt,
    "md": make_md,
    "docx": make_docx,
    "pdf": make_pdf,
    "html": make_html,
}


def build(dest_dir: str | Path, ext: str, name: str | None = None,
          content: str = KB_TEXT) -> Path:
    """在指定目录生成一个指定格式的文档。"""
    ext = ext.lower().lstrip(".")
    if ext not in FACTORY:
        raise ValueError(f"不支持的格式：{ext}，可选 {list(FACTORY)}")
    dest = Path(dest_dir) / (name or f"kb_sample_{ext}.{ext}")
    return FACTORY[ext](dest, content)


def build_kb_doc(ext: str = "txt", name: str | None = None) -> Path:
    """生成知识库文档到系统临时目录（用例间互不干扰）。"""
    tmp_dir = Path(tempfile.mkdtemp(prefix="docqa_kb_"))
    return build(tmp_dir, ext, name=name or KB_DOC_NAME, content=KB_TEXT)
